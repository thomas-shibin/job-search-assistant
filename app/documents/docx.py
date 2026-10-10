"""Deterministic DOCX rendering of a tailored resume from already-validated,
structured data (Milestone 6/7) -- the fallback used only when the base
resume is a PDF, so there is no Word document to keep. A .docx base resume is
tailored in place instead, keeping its own formatting (`app/documents/docx_edit.py`).

Callers (see `app/resume/tailor.py`) are responsible for already having
resolved which text to show for every section (validated AI text vs. a
verbatim fallback vs. dropped) -- this module only lays out whatever it's
given.
"""
from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

if TYPE_CHECKING:
    from app.config.schema import AddressProfile
    from app.models.candidate import Certification, Education, Employment


def _format_date_range(employment: "Employment") -> str:
    if employment.start_date is None:
        return ""
    start = employment.start_date.strftime("%b %Y")
    if employment.is_current or employment.end_date is None:
        end = "Present"
    else:
        end = employment.end_date.strftime("%b %Y")
    return f"{start} - {end}"


def _format_education_line(edu: "Education") -> str:
    degree_field = edu.degree or ""
    if edu.field_of_study:
        degree_field = f"{degree_field} in {edu.field_of_study}" if degree_field else edu.field_of_study
    return f"{degree_field}, {edu.institution}" if degree_field else edu.institution


def _format_certification_line(cert: "Certification") -> str:
    return f"{cert.name}, {cert.issuer}" if cert.issuer else cert.name


def render_tailored_resume(
    path: Path,
    *,
    candidate_name: str,
    contact_line: str | None,
    address: "AddressProfile | None",
    summary: str,
    skills_order: list[str],
    employment_bullets: list[tuple["Employment", list[str]]],
    education: list["Education"],
    certifications: list["Certification"],
) -> None:
    """Writes a resume DOCX to `path`. `employment_bullets` is already ordered
    (most recent employment first) and already contains only the final bullet
    text to display for each -- see `app/resume/tailor.py`.

    `contact_line` (email/phone/links, extracted from the base resume's own
    header by `app/resume/parser.py`) and `address` (an explicit city/
    province/country override swapped in only for a job tied to a search
    profile -- see `app/resume/tailor.py::resolve_address`) are independent:
    both render if both are given, rather than one trying to edit the other.
    """
    doc = Document()

    name_para = doc.add_paragraph()
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = name_para.add_run(candidate_name)
    name_run.bold = True
    name_run.font.size = Pt(18)

    if address is not None:
        addr_para = doc.add_paragraph(address.one_line())
        addr_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if contact_line:
        contact_para = doc.add_paragraph(contact_line)
        contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if summary:
        doc.add_heading("Summary", level=1)
        doc.add_paragraph(summary)

    if skills_order:
        doc.add_heading("Skills", level=1)
        doc.add_paragraph(", ".join(skills_order))

    if employment_bullets:
        doc.add_heading("Experience", level=1)
        for employment, bullet_texts in employment_bullets:
            title_para = doc.add_paragraph()
            title_para.add_run(f"{employment.title}, {employment.company}").bold = True
            date_range = _format_date_range(employment)
            if date_range:
                doc.add_paragraph(date_range)
            for bullet_text in bullet_texts:
                doc.add_paragraph(bullet_text, style="List Bullet")

    if education:
        doc.add_heading("Education", level=1)
        for edu in education:
            doc.add_paragraph(_format_education_line(edu))

    if certifications:
        doc.add_heading("Certifications", level=1)
        for cert in certifications:
            doc.add_paragraph(_format_certification_line(cert))

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
