"""Canadian resume check: what in a resume works against it with Canadian
employers, and the fixes that can be made without changing what it claims.

Deterministic and free (no AI): it reads the resume's own lines and finds
- a photo, and personal details Canadian employers don't ask for (date of
  birth, marital status, nationality, religion, family, ID numbers...) --
  human-rights guidance tells employers not to ask for them;
- a declaration ("I hereby declare...") and listed references;
- dates written as numbers ("03/2019", "05/04/2019"), which read differently
  in different countries;
- British or American spelling where Canadian spelling differs (colour,
  centre, travelled; but organize, analyze);
- high school listed next to a degree, unfamiliar grading, a missing city and
  province, a foreign phone number, an objective in place of a summary, and
  length.

Each finding's source and what it costs the job seeker are in
app/modules/resumes/check_sources.py; a rule with no official source says so.

`check()` turns these into `Finding`s for the page; `plan()` turns the ones the
member chose into line edits (`Plan`), which app/modules/resumes/canadian_convert.py
writes into their own Word file. Fixes only remove or reformat what's there --
except the Canadian equivalent of a job title, which the member supplies or
approves and which is added next to the original title, never in its place.
"""
from __future__ import annotations

import math
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document

from app.modules.resumes.check_sources import Basis, basis
from app.modules.resumes.documents.docx_edit import PHOTO_MIN_INCHES, iter_paragraphs, _photo_runs
from app.modules.resumes.parser import _extract_pdf_lines, _match_section_key

MUST, SHOULD, TIP = "must", "should", "tip"
LEVELS = (MUST, SHOULD, TIP)

# Fix keys (each `Finding.key` with a fix; the form posts the chosen ones).
PHOTO, PERSONAL, DECLARATION, REFERENCES, HIGH_SCHOOL, DATES, SPELLING = (
    "photo", "personal", "declaration", "references", "high_school", "dates", "spelling")
FIXES = (PHOTO, PERSONAL, DECLARATION, REFERENCES, HIGH_SCHOOL, DATES, SPELLING)

MAX_EXAMPLES = 12
WORDS_PER_PAGE = 600  # a dense resume page; only used when the file doesn't say


# ---- reading the resume ------------------------------------------------------------------

@dataclass
class Line:
    text: str
    heading: bool
    section: str  # "header" (name and contact), a section key (see _section_of), or "other"


@dataclass
class ResumeDocument:
    lines: list[Line]
    photos: int
    pages: int
    is_docx: bool


_OWN_SECTIONS = {
    "personal": ("personal details", "personal information", "personal profile", "personal data",
                 "personal particulars", "personal info", "personal"),
    "references": ("references", "referees", "reference", "professional references"),
    "declaration": ("declaration", "self declaration", "self-declaration"),
    "objective": ("objective", "career objective", "job objective", "professional objective"),
    "education": ("education", "educational qualifications", "educational qualification",
                  "academic qualifications", "academic qualification", "educational background",
                  "academic background", "academics", "education and training", "academic details",
                  "qualifications", "education & training"),
}


def _section_of(text: str) -> str | None:
    """The section a heading starts, or None when it isn't a known heading."""
    normalized = " ".join(text.strip().rstrip(":").split()).lower()
    for key, aliases in _OWN_SECTIONS.items():
        if normalized in aliases:
            return key
    return _match_section_key(normalized)


_DIGIT = re.compile(r"\d")


def _short_label(text: str) -> bool:
    """Short, no digits, no "label: value": could be a heading."""
    stripped = text.strip().rstrip(":")
    return 0 < len(stripped) <= 40 and not _DIGIT.search(stripped) and ":" not in stripped


def _sectioned(raw: list[tuple[str, bool, bool, bool]]) -> list[Line]:
    """(text, looks like a heading, bold-only line, in page header) -> lines with
    their section. Bold-only lines count as headings only in a resume whose known
    headings are themselves just bold (otherwise they're job titles, degrees...)."""
    bold_headings = any(bold and _section_of(text) for text, _, bold, _ in raw)
    lines, section, seen_body = [], "header", False
    for text, heading_hint, bold, in_header in raw:
        if in_header:
            lines.append(Line(text, False, "header"))
            continue
        known = _section_of(text)
        heading = bool(known) or (seen_body and _short_label(text) and (heading_hint or (bold and bold_headings)))
        seen_body = True  # the first line is the name, whatever its style
        if heading:
            section = known or "other"
        lines.append(Line(text, heading, section))
    return lines


def _docx_pages(path: Path, words: int) -> int:
    """Pages as Word last counted them (docProps/app.xml), or an estimate."""
    estimate = max(1, math.ceil(words / WORDS_PER_PAGE))
    try:
        with zipfile.ZipFile(path) as archive:
            match = re.search(rb"<Pages>(\d+)</Pages>", archive.read("docProps/app.xml"))
        return max(int(match.group(1)), estimate) if match else estimate
    except (KeyError, OSError, zipfile.BadZipFile):
        return estimate


def read_docx(path: Path) -> ResumeDocument:
    doc = Document(str(path))
    raw = []
    for paragraph, in_header in iter_paragraphs(doc):
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph._p.style or "").lower()
        runs = [r for r in paragraph.runs if r.text.strip()]
        bold = bool(runs) and all(r.bold for r in runs) and "list" not in style
        caps = text.isupper() and any(c.isalpha() for c in text)
        raw.append((text, style.startswith("heading") or (caps and "list" not in style), bold, in_header))
    lines = _sectioned(raw)
    words = sum(len(line.text.split()) for line in lines)
    return ResumeDocument(lines, len(_photo_runs(doc, PHOTO_MIN_INCHES)), _docx_pages(path, words), True)


_PHOTO_MIN_POINTS = PHOTO_MIN_INCHES * 72


def read_pdf(path: Path) -> ResumeDocument:
    import pdfplumber

    raw = [(info.text, info.is_heading_hint, False, False) for info in _extract_pdf_lines(path)]
    with pdfplumber.open(str(path)) as pdf:
        pages = len(pdf.pages)
        photos = sum(1 for image in (pdf.pages[0].images if pdf.pages else [])
                     if min(float(image.get("width", 0)), float(image.get("height", 0))) >= _PHOTO_MIN_POINTS)
    return ResumeDocument(_sectioned(raw), photos, max(pages, 1), False)


def read(path: Path) -> ResumeDocument:
    return read_docx(path) if path.suffix.lower() == ".docx" else read_pdf(path)


# ---- personal details ----------------------------------------------------------------------

_PERSONAL: list[tuple[str, re.Pattern]] = [(label, re.compile(pattern)) for label, pattern in [
    ("date of birth", r"(?i:\bdate\s+of\s+birth\b|\bbirth\s*date\b|\bd\.?\s?o\.?\s?b\b|\bborn\s+(?:on|in)\b)"),
    ("age", r"(?i:\bage\s*[:\-–]\s*\d{2}\b|\b\d{2}\s+years\s+old\b)"),
    ("gender", r"(?i:\b(?:gender|sex)\s*[:\-–])"),
    ("marital status", r"(?i:\bmarital\s+status\b|\bmarried\b|\bunmarried\b)"),
    ("nationality", r"(?i:\bnationality\s*[:\-–]|\bcitizenship\s*[:\-–](?!\s*canad))"),
    ("religion", r"(?i:\b(?:religion|caste)\s*[:\-–])"),
    ("family details", r"(?i:\b(?:father|mother|husband|wife|spouse|guardian)(?:['’]s)?\s+(?:name|occupation)\b"
                       r"|\b(?:number\s+of\s+)?(?:children|dependants|dependents)\s*[:\-–])"),
    ("ID numbers", r"(?i:\bpassport\s*(?:no\.?|number|#))|(?i:\bsocial\s+insurance\s+number\b)|\bSIN\s*(?:[:#]|[Nn]o\b)"
                   r"|(?i:\baadhaa?r\b)|(?i:\bnational\s+id(?:entity)?\s*(?:card|no\.?|number)?\s*[:#])"
                   r"|\bPAN\s*(?:[Nn]o\b|[Nn]umber|[Cc]ard)"),
    ("height, weight or health", r"(?i:\b(?:height|weight|blood\s+group|health(?:\s+status)?)\s*[:\-–])"),
    ("salary", r"(?i:\b(?:current|expected|present|desired)\s+(?:salary|compensation|package|pay)\b)|\bCTC\b"),
]]
_BARE_PERSONAL = {"male", "female", "married", "single", "unmarried", "divorced", "widowed"}
_SEPARATOR = re.compile(r"(\s*(?:\||•|·|\t|;|\s{3,})\s*)")


def _personal_kinds(text: str) -> list[str]:
    return [label for label, pattern in _PERSONAL if pattern.search(text)]


def without_personal(text: str) -> tuple[str | None, list[str]]:
    """`text` with its personal-detail parts taken out (None: nothing left), and
    which kinds of detail they were. Parts are split at |, •, tabs and the like,
    so "jordan@example.com | DOB: 1990 | Victoria, BC" keeps its other parts."""
    kinds = _personal_kinds(text)
    if not kinds:
        return text, []
    parts = _SEPARATOR.split(text)
    segments, separators = parts[0::2], parts[1::2]
    keep = [i for i, segment in enumerate(segments) if segment.strip() and not _personal_kinds(segment)]
    if not keep:
        return None, kinds
    out = segments[keep[0]]
    for i in keep[1:]:
        out += separators[i - 1] + segments[i]
    return out.strip(), kinds


def _personal_edits(lines: list[Line]) -> tuple[dict[int, str | None], list[str]]:
    edits: dict[int, str | None] = {}
    kinds: list[str] = []
    for i, line in enumerate(lines):
        if line.heading:
            continue
        if line.section == "personal" and line.text.strip().strip(".").lower() in _BARE_PERSONAL:
            edits[i] = None
            kinds.append("marital status or gender")
            continue
        new, found = without_personal(line.text)
        if found:
            edits[i] = new
            kinds.extend(found)
    # A "Personal details" heading goes too once nothing is left under it.
    for i, line in enumerate(lines):
        if line.heading and line.section == "personal":
            under = []
            for j in range(i + 1, len(lines)):
                if lines[j].heading:
                    break
                under.append(j)
            if under and all(edits.get(j, "") is None for j in under):
                edits[i] = None
    return edits, list(dict.fromkeys(kinds))


# ---- declaration, references, high school --------------------------------------------------

_DECLARE = re.compile(r"(?i)\bhereby\s+(?:declare|certify|affirm|state)\b"
                      r"|\bthe\s+above\s+(?:information|particulars|details|statements?)\s+(?:is|are)\s+(?:true|correct)")
_SIGN_OFF = re.compile(r"(?i)^\s*(?:place|date|signature|signed|yours\s+(?:faithfully|sincerely))\b")
_REFERENCES_ON_REQUEST = re.compile(
    r"(?i)\b(?:references?|referees?)\b.{0,40}\b(?:available|provided|furnished|given)\b.{0,20}\b(?:request|demand|asked)\b")


def _section_edits(lines: list[Line], key: str) -> dict[int, None]:
    """Every line of each `key` section, heading included."""
    return {i: None for i, line in enumerate(lines) if line.section == key}


def _declaration_edits(lines: list[Line]) -> dict[int, None]:
    edits = _section_edits(lines, "declaration")
    for i, line in enumerate(lines):
        if _DECLARE.search(line.text):
            edits[i] = None
            for j in range(i + 1, min(i + 4, len(lines))):  # "Place: ...", "Date: ...", "Signature"
                if _SIGN_OFF.match(lines[j].text) and not lines[j].heading:
                    edits[j] = None
    return edits


def _reference_edits(lines: list[Line]) -> tuple[dict[int, None], bool]:
    """(edits, whether referees' details are listed)."""
    edits = _section_edits(lines, "references")
    listed = any(not lines[i].heading and not _REFERENCES_ON_REQUEST.search(lines[i].text) for i in edits)
    for i, line in enumerate(lines):
        if _REFERENCES_ON_REQUEST.search(line.text):
            edits[i] = None
    return edits, listed


_HIGH_SCHOOL = re.compile(
    r"(?i:\b(?:ssc|hsc|ssce|sslc|cbse|icse|igcse|waec|higher\s+secondary|senior\s+secondary|secondary\s+school"
    r"|high\s+school|matriculation|matric|10th|12th|class\s+(?:x|xii|10|12)|grade\s+12)\b)|\b[AO][-\s]Levels?\b")
_JUNIOR_COLLEGE = re.compile(r"(?i)\bjunior\s+college\b")  # high school in India
_SCHOOL_DETAIL = re.compile(r"(?i)\b(?:school|vidyalaya|academy|board|lyc[ée]e|colegio)\b")


_DEGREE = re.compile(
    r"(?i:\b(?:bachelor|master|doctor|diploma|degree|certificate|university|college|institute|polytechnic"
    r"|engineering)\b)|\b(?:B\.?\s?(?:A|S|Sc|Com|Tech|E|Eng|BA|Ed)|M\.?\s?(?:A|S|Sc|Com|Tech|E|Eng|BA|Ed)|MBA|PhD|BBA|BCA|MCA)\b")


def is_high_school(text: str) -> bool:
    return bool(_HIGH_SCHOOL.search(text))


def _high_school_edits(lines: list[Line]) -> dict[int, None]:
    """High-school entries in Education -- only when a post-secondary one stays."""
    entries = [i for i, line in enumerate(lines) if line.section == "education" and not line.heading]
    marked = [i for i in entries if _HIGH_SCHOOL.search(lines[i].text)]
    edits: dict[int, None] = {}
    for i in marked:
        edits[i] = None
        following = i + 1
        detail = _JUNIOR_COLLEGE.sub("school", lines[following].text) if following in entries else ""
        if following not in marked and _SCHOOL_DETAIL.search(detail) and not _DEGREE.search(detail):
            edits[following] = None  # "ABC School, 2008, 85%" under "Higher Secondary Certificate"
    if not any(_DEGREE.search(lines[i].text) for i in entries if i not in edits):
        return {}
    return edits


# ---- dates -----------------------------------------------------------------------------------

_MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
_FULL_MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September",
                "October", "November", "December"]
_FULL_DATE = re.compile(r"(?<![\d/.\-])(\d{1,2})\s*([/.\-])\s*(\d{1,2})\s*\2\s*((?:19|20)\d{2})(?!\d)")
_MONTH_YEAR = re.compile(r"(?<![\d/.\-])(0?[1-9]|1[0-2])\s*[/.\-]\s*((?:19|20)\d{2})(?![\d])(?![/.\-]\d)")
_FULL_NAMED = re.compile(r"\b(?:" + "|".join(_FULL_MONTHS) + r")\s+(?:19|20)\d{2}\b")
_SHORT_NAMED = re.compile(r"\b(?:" + "|".join(_MONTHS) + r")\.?\s+(?:19|20)\d{2}\b")


def month_style(lines: list[Line]) -> list[str]:
    """Month names the way the resume already writes them ("January" or "Jan")."""
    text = "\n".join(line.text for line in lines)
    return _FULL_MONTHS if len(_FULL_NAMED.findall(text)) > len(_SHORT_NAMED.findall(text)) else _MONTHS


def redate(text: str, names: list[str] = _MONTHS) -> tuple[str, list[tuple[str, str]], list[str]]:
    """Numeric dates as "Mar 2019": (new text, [(old, new)], dates whose day and
    month can't be told apart, left as written)."""
    changed: list[tuple[str, str]] = []
    unclear: list[str] = []

    def full(match: re.Match) -> str:
        a, b, year = int(match.group(1)), int(match.group(3)), match.group(4)
        if a > 12 and 1 <= b <= 12:
            month = b  # day first: 25/03/2019
        elif b > 12 and 1 <= a <= 12 or a == b and 1 <= a <= 12:
            month = a  # month first: 03/25/2019
        else:
            if a <= 12 and b <= 12:
                unclear.append(match.group(0))
            return match.group(0)
        new = f"{names[month - 1]} {year}"
        changed.append((match.group(0), new))
        return new

    def month_year(match: re.Match) -> str:
        new = f"{names[int(match.group(1)) - 1]} {match.group(2)}"
        changed.append((match.group(0), new))
        return new

    text = _FULL_DATE.sub(full, text)
    text = _MONTH_YEAR.sub(month_year, text)
    return text, changed, unclear


# ---- spelling --------------------------------------------------------------------------------

def _spellings() -> dict[str, str]:
    words: dict[str, str] = {}
    our = {"color": ("", "s", "ed", "ing", "ful"), "behavior": ("", "s", "al", "ally"),
           "labor": ("", "s", "ed", "ing", "er", "ers"), "favor": ("", "s", "ed", "ing", "able", "ite", "ites"),
           "honor": ("", "s", "ed", "ing", "able"), "neighbor": ("", "s", "ing", "hood", "hoods"),
           "endeavor": ("", "s", "ed", "ing"), "flavor": ("", "s", "ed", "ing"), "harbor": ("", "s", "ed", "ing"),
           "rumor": ("", "s", "ed"), "savor": ("", "s", "ed", "ing"), "armor": ("", "ed"), "odor": ("", "s"),
           "vigor": ("",), "valor": ("",), "candor": ("",), "parlor": ("", "s")}
    for base, endings in our.items():
        for ending in endings:
            words[base + ending] = base[:-2] + "our" + ending
    for base in ("center", "theater", "liter", "fiber", "caliber", "somber", "meager", "luster"):
        words[base] = base[:-2] + "re"
        words[base + "s"] = base[:-2] + "res"
    words.update({"centered": "centred", "centering": "centring", "centerpiece": "centrepiece"})
    for us, ca in {"travel": "travell", "cancel": "cancell", "model": "modell", "label": "labell",
                   "counsel": "counsell", "fuel": "fuell", "level": "levell", "signal": "signall",
                   "total": "totall", "channel": "channell", "marvel": "marvell"}.items():
        words[us + "ed"], words[us + "ing"] = ca + "ed", ca + "ing"
    words.update({"traveler": "traveller", "travelers": "travellers", "counselor": "counsellor",
                  "counselors": "counsellors", "jewelry": "jewellery", "catalog": "catalogue",
                  "catalogs": "catalogues", "defense": "defence", "offense": "offence", "gray": "grey",
                  "enrollment": "enrolment", "enrollments": "enrolments", "acknowledgment": "acknowledgement",
                  "acknowledgments": "acknowledgements"})
    # British -ise / -yse: Canadian spelling uses -ize / -yze.
    for stem in ("organ", "recogn", "real", "special", "optim", "priorit", "util", "final", "standard", "custom",
                 "minim", "maxim", "summar", "author", "categor", "mobil", "modern", "central", "capital", "visual",
                 "synchron", "initial", "normal", "stabil", "apolog", "critic", "character", "familiar", "memor",
                 "monet", "personal", "public", "sanit", "systemat", "emphas", "digit", "digital", "global", "local",
                 "general", "human", "industrial", "commercial", "energ", "jeopard", "legal", "neutral", "penal",
                 "rational", "revolution", "subsid", "symbol", "synthes", "harmon", "hospital", "computer",
                 "democrat", "fertil", "galvan", "hypothes", "immun", "incentiv", "item", "legitim", "operational",
                 "conceptual", "contextual", "decentral", "commodit", "strateg"):
        for ending in ("e", "es", "ed", "ing", "ation", "ations", "er", "ers", "able"):
            words[f"{stem}is{ending}"] = f"{stem}iz{ending}"
    for stem in ("analy", "cataly", "paraly"):
        for ending in ("se", "sed", "sing", "ser", "sers"):
            words[stem + ending] = stem + "z" + ending[1:]
    return words


_SPELLING = _spellings()
_WORD = re.compile(r"[A-Za-z]+")
_SENTENCE_START = re.compile(r"(?:^|[.!?:;]\s+|^[\s•●○◦▪■‣⁃∙·\-–—*]+)$")


def _cased(word: str, replacement: str) -> str:
    if word.isupper():
        return replacement.upper()
    return replacement[:1].upper() + replacement[1:] if word[:1].isupper() else replacement


def respell(text: str) -> tuple[str, list[tuple[str, str]]]:
    """Canadian spelling: (new text, [(old word, new word)]). A capitalized word
    mid-sentence is a name ("Kennedy Space Center", "Labor Day") and kept, as are
    words in all capitals, web addresses and emails."""
    changed: list[tuple[str, str]] = []
    out, last = [], 0
    for match in _WORD.finditer(text):
        word = match.group()
        replacement = _SPELLING.get(word.lower())
        if replacement is None or word.isupper() and len(word) > 1:
            continue
        start, end = match.span()
        before, after, then = text[start - 1:start], text[end:end + 1], text[end + 1:end + 2]
        if before in ("/", "@", ".") or after in ("/", "@") or (after == "." and then.isalnum()):
            continue  # part of a web address or email
        if word[:1].isupper() and not _SENTENCE_START.search(text[:start]):
            continue
        new = _cased(word, replacement)
        out.append(text[last:match.start()] + new)
        last = match.end()
        changed.append((word, new))
    out.append(text[last:])
    return "".join(out), changed


# ---- the check ---------------------------------------------------------------------------------

@dataclass
class Finding:
    key: str
    level: str  # MUST | SHOULD | TIP
    title: str
    advice: str
    examples: list[str] = field(default_factory=list)
    fix_label: str | None = None  # what fixing it does; None: advice only
    unclear: list[str] = field(default_factory=list)  # parts the fix leaves for the member
    lines: list[int] = field(default_factory=list)  # the resume lines it's about (indexes into doc.lines)

    @property
    def fixable(self) -> bool:
        return self.fix_label is not None

    @property
    def basis(self) -> Basis | None:
        """Where the rule comes from and what the problem costs (check_sources.py)."""
        return basis(self.key)


_PROVINCE = re.compile(r"\b(?:AB|BC|MB|NB|NL|NS|NT|NU|ON|PE|PEI|QC|SK|YT)\b|(?i:\b(?:canada|alberta|british\s+columbia"
                       r"|manitoba|new\s+brunswick|newfoundland|nova\s+scotia|northwest\s+territories|nunavut"
                       r"|ontario|prince\s+edward\s+island|qu[ée]bec|saskatchewan|yukon)\b)")
_FOREIGN_PHONE = re.compile(r"\+(?!1\b|1[\s(.-])\d{1,3}[\s(.-]*\d")
_GRADES = re.compile(r"(?i)\b(?:c?gpa|cpi|sgpa)\b[^.\n]{0,15}\d+(?:\.\d+)?\s*/\s*10\b|\b\d{2}(?:\.\d+)?\s*%|\bpercentage\b"
                     r"|\b(?:first|second)\s+class\b|\bdistinction\b")
_FOREIGN_DEGREE = re.compile(r"\b(?:B\.?\s?Tech|M\.?\s?Tech|B\.\s?E\.|B\.?\s?Com|M\.?\s?Com|BCA|MCA|PGDM|PGDBA|MBBS|"
                             r"B\.?\s?Pharm|LL\.?B|Licenciatura|Licenciado|Ingenier[oa]|Bacharelado|Diplom-)\b")
_OBJECTIVE_WORDS = re.compile(r"(?i)\b(?:seeking|to\s+obtain|to\s+secure|looking\s+for|to\s+work\s+in)\b.{0,60}"
                              r"\b(?:challenging|position|role|opportunity|organi[sz]ation)\b")


def _listed(items: list[str]) -> str:
    """"a", "a and b", "a, b and c"."""
    return items[0] if len(items) == 1 else f"{', '.join(items[:-1])} and {items[-1]}"


def _examples(values: list[str]) -> list[str]:
    values = list(dict.fromkeys(values))
    return values[:MAX_EXAMPLES] + ([f"and {len(values) - MAX_EXAMPLES} more"] if len(values) > MAX_EXAMPLES else [])


def check(doc: ResumeDocument) -> list[Finding]:
    """What to change for Canadian employers, most important first."""
    lines = doc.lines
    findings: list[Finding] = []

    if doc.photos:
        findings.append(Finding(
            PHOTO, MUST, "Remove the photo",
            "Canadian resumes don't include a photo. Employers avoid anything that could suggest bias "
            "based on appearance, age or ethnicity.",
            fix_label="Remove the photo"))

    edits, kinds = _personal_edits(lines)
    if edits:
        findings.append(Finding(
            PERSONAL, MUST, "Remove personal details",
            f"Leave out your {_listed(kinds)}. Canadian human-rights law keeps hiring from depending on "
            "these, so employers don't ask for them, and a resume that includes them looks out of place. "
            "Your name, city, phone, email and LinkedIn are enough.",
            _examples([lines[i].text for i in edits if not lines[i].heading]),
            fix_label="Remove these details", lines=sorted(edits)))

    edits = _declaration_edits(lines)
    if edits:
        findings.append(Finding(
            DECLARATION, SHOULD, "Remove the declaration",
            "A signed declaration (\"I hereby declare that the above information is true\"), with place, date "
            "and signature, isn't used in Canada. Submitting the resume already says it's accurate.",
            _examples([lines[i].text for i in edits]), fix_label="Remove the declaration", lines=sorted(edits)))

    edits, listed = _reference_edits(lines)
    if edits:
        findings.append(Finding(
            REFERENCES, SHOULD if listed else TIP,
            "Take your references off the resume" if listed else "Drop \"References available upon request\"",
            ("Keep your referees' names and contact details for when an employer asks, usually after an "
             "interview. That protects their privacy and saves room.") if listed else
            "Employers in Canada assume you'll provide references when asked, so the line only takes up room.",
            _examples([lines[i].text for i in edits]), fix_label="Remove them", lines=sorted(edits)))

    edits = _high_school_edits(lines)
    if edits:
        findings.append(Finding(
            HIGH_SCHOOL, SHOULD, "Leave out high school",
            "Once you have a college or university credential, Canadian resumes list only that. High-school "
            "results take room without telling employers anything new.",
            _examples([lines[i].text for i in edits]), fix_label="Remove high school", lines=sorted(edits)))

    leaving = {i for i, value in _personal_edits(lines)[0].items() if value is None}
    leaving |= set(_declaration_edits(lines)) | set(_reference_edits(lines)[0]) | set(_high_school_edits(lines))
    names = month_style(lines)
    changed, unclear, dated = [], [], []
    for i, line in enumerate(lines):
        if line.section == "header" or i in leaving:
            continue
        _, line_changed, line_unclear = redate(line.text, names)
        changed += [f"{old} → {new}" for old, new in line_changed]
        unclear += line_unclear
        if line_changed or line_unclear:
            dated.append(i)
    if changed or unclear:
        findings.append(Finding(
            DATES, SHOULD, "Write dates as month and year",
            "Dates in numbers read differently from country to country (is 05/04/2019 April or May?). "
            f"Canadian resumes use the month's name and the year, like \"{names[2]} 2019\".",
            _examples(changed), fix_label="Change the dates" if changed else None,
            unclear=list(dict.fromkeys(unclear)), lines=dated))

    spelled, respelled = [], []
    for i, line in enumerate(lines):
        if line.section != "header" and i not in leaving:
            words = respell(line.text)[1]
            spelled += [f"{old} → {new}" for old, new in words]
            if words:
                respelled.append(i)
    if spelled:
        # A tip, not "should fix": the Language Portal of Canada says there's no single Canadian standard.
        findings.append(Finding(
            SPELLING, TIP, "Use Canadian spelling",
            "Canadian spelling mixes British and American: colour, centre, travelled and licence, "
            "but organize, analyze and program. Employers notice, especially for writing-heavy roles.",
            _examples(spelled), fix_label="Change the spelling", lines=respelled))

    if doc.pages >= 3:
        findings.append(Finding(
            "length", MUST if doc.pages >= 4 else SHOULD, f"Shorten it to two pages (it's about {doc.pages})",
            "Most Canadian employers expect one or two pages; three is fine only for long senior or academic "
            "careers. Keep the last 10 to 15 years in detail, sum up earlier roles in a line or two, and cut "
            "points that don't relate to the jobs you want. Tailoring for a job lets you drop points per application."))

    header = [line.text for line in lines if line.section == "header"]
    header_at = [i for i, line in enumerate(lines) if line.section == "header"]
    if header and not any(_PROVINCE.search(text) for text in header):
        findings.append(Finding(
            "location", SHOULD, "Add your city and province",
            "Employers and their screening software look for where you are. Put your city and province under "
            "your name (\"Toronto, ON\"). If you're moving, say so: \"Relocating to Calgary, AB in May 2027\"."))
    # No street-address finding: Job Bank's own advice is to list your address.
    phones = [i for i in header_at if _FOREIGN_PHONE.search(lines[i].text)]
    if phones:
        findings.append(Finding(
            "phone", TIP, "Add a Canadian phone number",
            "Employers rarely call international numbers. If you have a Canadian number, list it first; if "
            "not, a free Canadian number you can answer helps.", _examples([lines[i].text for i in phones]),
            lines=phones))

    education_at = [i for i, line in enumerate(lines)
                    if line.section == "education" and not line.heading and i not in leaving]
    grades = [lines[i].text for i in education_at if _GRADES.search(lines[i].text)]
    foreign = [lines[i].text for i in education_at if _FOREIGN_DEGREE.search(lines[i].text)]
    if grades or foreign:
        findings.append(Finding(
            "credentials", TIP, "Help employers understand your education",
            "Canadian employers may not know degrees from other countries or grades out of 10 or in percent. "
            "If you've had your credentials assessed (for example by WES), add the result, such as \"Bachelor's "
            "degree (WES: equivalent to a Canadian four-year bachelor's)\". Leave grades out unless they're a "
            "real strength.", _examples(foreign + grades),
            lines=[i for i in education_at if _GRADES.search(lines[i].text) or _FOREIGN_DEGREE.search(lines[i].text)]))

    objective = [i for i, line in enumerate(lines) if line.heading and line.section == "objective"
                 or line.section in ("objective", "summary") and _OBJECTIVE_WORDS.search(line.text)]
    if objective:
        findings.append(Finding(
            "objective", TIP, "Replace the objective with a summary",
            "Instead of what you're looking for (\"seeking a challenging position...\"), write two or three lines "
            "on what you offer: your field, years of experience and strongest skills. You can do this on the "
            "Edit resume page.", lines=objective))

    order = {level: n for n, level in enumerate(LEVELS)}
    return sorted(findings, key=lambda f: order[f.level])


# ---- the fixes -------------------------------------------------------------------------------------

@dataclass
class Plan:
    edits: dict[int, str | None]  # line index -> its new text, or None to take it out
    photos: bool = False
    done: list[str] = field(default_factory=list)  # what the fixes did, for the member
    titles_missed: list[str] = field(default_factory=list)


def _count(n: int, word: str) -> str:
    return f"{n} {word}{'s' if n != 1 else ''}"


def plan(doc: ResumeDocument, chosen: set[str], titles: dict[str, str] | None = None,
         only: dict[str, set[int]] | None = None) -> Plan:
    """The edits for the fixes in `chosen` (keys from FIXES), plus each job title
    in `titles` followed by its Canadian equivalent: "Executive - HR (HR Coordinator)".
    `only` limits a fix to some lines ({key: line indexes}; the photo is line -1),
    for a member who picked the changes one by one; a key not in it changes every line."""
    lines = doc.lines
    result = Plan(edits={})
    text = {i: line.text for i, line in enumerate(lines)}
    only = only or {}

    def allowed(key: str, i: int) -> bool:
        return key in chosen and (key not in only or i in only[key])

    def apply(key: str, edits: dict[int, str | None]) -> int:
        n = 0
        for i, value in edits.items():
            if text.get(i) is not None and allowed(key, i):
                text[i] = value
                n += 1
        return n

    if allowed(PHOTO, -1) and doc.photos:
        result.photos = True
        result.done.append("removed the photo")
    if (n := apply(PERSONAL, _personal_edits(lines)[0])):
        result.done.append(f"removed personal details from {_count(n, 'line')}")
    if apply(DECLARATION, _declaration_edits(lines)):
        result.done.append("removed the declaration")
    if apply(REFERENCES, _reference_edits(lines)[0]):
        result.done.append("removed the references")
    if apply(HIGH_SCHOOL, _high_school_edits(lines)):
        result.done.append("removed high school")

    names = month_style(lines)
    dates = words = 0
    for i, line in enumerate(lines):
        if text[i] is None or line.section == "header":
            continue
        if allowed(DATES, i):
            text[i], changed, _ = redate(text[i], names)
            dates += len(changed)
        if allowed(SPELLING, i):
            text[i], changed = respell(text[i])
            words += len(changed)
    if dates:
        result.done.append(f"rewrote {_count(dates, 'date')}")
    if words:
        result.done.append(f"changed the spelling of {_count(words, 'word')}")

    added = 0
    for original, equivalent in (titles or {}).items():
        spot = next((i for i, line in enumerate(lines)
                     if text[i] is not None and line.section != "header" and not line.heading
                     and original in text[i] and f"{original} (" not in text[i]), None)
        if spot is None:
            result.titles_missed.append(original)
            continue
        text[spot] = text[spot].replace(original, f"{original} ({equivalent})", 1)
        added += 1
    if added:
        result.done.append(f"added the Canadian equivalent to {_count(added, 'job title')}")

    result.edits = {i: value for i, value in text.items() if value != lines[i].text}
    return result


# ---- each change, for choosing them one by one -----------------------------------------------------

TOP = "Top of the page"


@dataclass
class Change:
    key: str  # the fix it belongs to
    line: int  # its line index (-1: the photo)
    before: str
    after: str | None  # None: taken out
    section: str  # the heading it's under, as the resume words it


def headings(doc: ResumeDocument) -> list[str]:
    """For each line, the heading it sits under ("Work Experience"), or TOP."""
    out, current = [], TOP
    for line in doc.lines:
        if line.heading and line.section != "header":
            current = line.text.strip().rstrip(":")
        out.append(TOP if line.section == "header" else current)
    return out


def changes(doc: ResumeDocument) -> dict[str, list[Change]]:
    """Every change each fix would make, line by line: what the page shows
    before anything is changed, so the member can choose."""
    under = headings(doc)
    out: dict[str, list[Change]] = {}
    if doc.photos:
        out[PHOTO] = [Change(PHOTO, -1, "Your photo", None, TOP)]
    for key in FIXES:
        if key == PHOTO:
            continue
        edits = plan(doc, {key}).edits
        if edits:
            out[key] = [Change(key, i, doc.lines[i].text, new, under[i]) for i, new in sorted(edits.items())]
    return out


# ---- is it a resume at all? ---------------------------------------------------------------------
# The checks look for problems, so a file with no resume in it (a health card, a
# letter, a form) would have none and score as ready. It's checked first instead.

CORE_SECTIONS = {"summary", "objective", "experience", "education", "skills", "certifications", "projects"}
_RANGE_END = r"(?:(?:19|20)\d{2}|present|current|now|today|date)"
_RANGE = re.compile(rf"(?i)\b(?:[a-z]{{3,9}}\.?\s+)?(?:\d{{1,2}}/)?(?:19|20)\d{{2}}\s*(?:[-–—]|to)\s*"
                    rf"(?:[a-z]{{3,9}}\.?\s+)?(?:\d{{1,2}}/)?{_RANGE_END}\b")
_EMAIL_TEXT = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
_PHONE_TEXT = re.compile(r"(?:\+?\d[\s().-]*){10,15}")
MIN_RESUME_WORDS = 30
# Headings a job posting has and a resume doesn't: a posting also has Skills and
# Qualifications, but never date ranges for jobs held.
_POSTING = re.compile(r"(?i)^(?:key\s+)?(?:responsibilities|requirements|what you['’]ll do|what you will do|"
                      r"about the (?:role|job|position|team|company)|about us|who you are|what we offer|"
                      r"what you['’]ll bring|benefits|perks|how to apply|nice to have|must have|"
                      r"job description|job summary|duties)\s*:?$")

NOT_A_RESUME = ("This doesn't look like a resume, so it wasn't checked. A resume has parts like Experience, "
                "Education or Skills, with dates for each job, and that's what the check looks for. If this is "
                "your resume, give each part its own heading line (for example \"Work Experience\" or "
                "\"Education\") and try again.")


@dataclass
class Evidence:
    words: int
    sections: set[str]  # the core resume sections it has a heading for
    ranges: int  # date ranges, like "Jan 2020 - Present" or "2016 – 2019"
    contact: bool  # an email address or phone number
    posting: int = 0  # job-posting headings ("Responsibilities", "How to apply")

    @property
    def is_resume(self) -> bool:
        if self.words < MIN_RESUME_WORDS or (self.posting >= 2 and self.ranges == 0):
            return False
        core = len(self.sections)
        return core >= 2 or (core == 1 and (self.ranges >= 1 or self.contact)) or (self.ranges >= 2 and self.contact)


def evidence(doc: ResumeDocument) -> Evidence:
    text = "\n".join(line.text for line in doc.lines)
    sections = {line.section for line in doc.lines if line.heading and line.section in CORE_SECTIONS}
    posting = sum(1 for line in doc.lines if _POSTING.match(line.text.strip()))
    return Evidence(len(text.split()), sections, len(_RANGE.findall(text)),
                    bool(_EMAIL_TEXT.search(text) or _PHONE_TEXT.search(text)), posting)
