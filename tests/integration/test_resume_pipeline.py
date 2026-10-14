"""Tests for `app/modules/resumes/pipeline.py::generate_resume_bundle` -- the shared
tailor -> PDF -> artifact bundle -> READY_FOR_REVIEW flow used by both the
CLI's `generate-resume` and the scheduler's automated run.
"""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.core.config.schema import AppConfig, CandidateConfig
from app.core.database.base import Base
from app.integrations.ai.mock import MockAIProvider
from app.integrations.libreoffice import PdfConversionError
import app.models  # noqa: F401  register all models on Base.metadata
from app.modules.jobs.models import Job, JobStatus
from app.modules.resumes.models import AccomplishmentBullet, CandidateProfile, Employment, Skill
from app.modules.resumes.pipeline import generate_resume_bundle


@pytest.fixture()
def session(tmp_path, monkeypatch) -> Session:
    monkeypatch.setenv("APPLICATIONS_DATA_DIR", str(tmp_path / "applications"))
    engine = create_engine(f"sqlite:///{tmp_path / 'test.db'}")
    Base.metadata.create_all(engine)
    with Session(engine) as s:
        yield s


def _seed(session: Session) -> Job:
    profile = CandidateProfile(raw_resume_path="x.pdf", source_hash="h", summary="Engineer.", is_active=True)
    profile.skills = [Skill(name="Java")]
    profile.employments = [
        Employment(
            company="Acme", title="Engineer", start_date=date(2020, 1, 1), is_current=True,
            bullets=[AccomplishmentBullet(text="Did the thing", order_index=0)],
        )
    ]
    session.add(profile)
    job = Job(
        source="manual", source_url="https://example.com", company="Acme", title="Engineer",
        description="d", status=JobStatus.QUALIFIED, fingerprint="fp", description_hash="dh",
    )
    session.add(job)
    session.commit()
    session.refresh(job)
    return job


def test_generate_resume_bundle_ends_with_ready_for_review_and_matching_snapshot(session, monkeypatch):
    # No LibreOffice in the test environment -- PDF conversion degrades gracefully.
    monkeypatch.setattr(
        "app.modules.resumes.pipeline.convert_to_pdf",
        lambda path: (_ for _ in ()).throw(PdfConversionError("no soffice")),
    )

    job = _seed(session)
    result = generate_resume_bundle(session, job, AppConfig(candidate=CandidateConfig(name="Jordan", base_resume="x.docx")), ai_provider=MockAIProvider())

    assert job.status == JobStatus.READY_FOR_REVIEW
    assert result.pdf_path is None
    assert result.pdf_error is not None

    job_json = json.loads(Path(result.artifact_set.job_json_path).read_text())
    assert job_json["status"] == "READY_FOR_REVIEW"  # snapshot reflects the final status, not RESUME_GENERATED


def test_tailored_resume_is_the_candidates_own_docx_with_only_wording_changed(session, tmp_path):
    from docx import Document
    from docx.shared import Pt, RGBColor

    from app.modules.resumes.tailor import tailor_resume

    base = tmp_path / "my_resume.docx"
    doc = Document()
    doc.add_paragraph().add_run("Jordan Casey").font.size = Pt(21)
    doc.add_paragraph("EXPERIENCE")
    bullet = doc.add_paragraph(style="List Bullet")
    run = bullet.add_run("Built the payment service that handles card refunds for customers")
    run.font.name, run.font.color.rgb = "Georgia", RGBColor(0x2F, 0x4F, 0x6F)
    doc.save(str(base))

    job = _seed(session)
    profile = session.query(CandidateProfile).one()
    profile.raw_resume_path = str(base)
    session.commit()

    seen = []

    class _Provider(MockAIProvider):
        def rewrite_passages(self, passages, job, requirements, candidate_skills):
            seen.extend(passages)
            return {p["id"]: "Built the customer payment service handling card refunds" for p in passages}

    result = tailor_resume(session, job, AppConfig(candidate=CandidateConfig(name="Jordan", base_resume=str(base))),
                           ai_provider=_Provider())

    # Only the sentence is sent for rewriting -- never the name or the heading.
    assert [(p["section"], p["text"]) for p in seen] == [
        ("EXPERIENCE", "Built the payment service that handles card refunds for customers")]
    out = Document(str(result.docx_path))
    texts = [p.text for p in out.paragraphs if p.text]
    assert texts == ["Jordan Casey", "EXPERIENCE", "Built the customer payment service handling card refunds"]
    styled = out.paragraphs[2]
    assert styled.style.name == "List Bullet"
    assert (styled.runs[0].font.name, styled.runs[0].font.color.rgb) == ("Georgia", RGBColor(0x2F, 0x4F, 0x6F))
    assert out.paragraphs[0].runs[0].font.size == Pt(21)


def _accept_all(draft, candidate):
    from app.modules.resumes.validator import BulletValidationResult, ValidationResult

    return ValidationResult(
        summary_accepted=True, skills_accepted=list(draft.skills_order), skills_rejected=[],
        bullets=[BulletValidationResult(text=b.text, source_bullet_id=b.source_bullet_id, accepted=True)
                 for b in draft.bullets],
    )


def test_rebuild_rewrites_an_existing_tailored_resume_from_saved_wording(session, tmp_path):
    from docx import Document

    from app.modules.resumes.tailoring_models import TailoredBulletProvenance, TailoredResume
    from app.modules.resumes.tailor import rebuild_tailored_resume

    base = tmp_path / "my_resume.docx"
    doc = Document()
    doc.add_paragraph("Engineer.")
    doc.add_paragraph("Did the thing", style="List Bullet")
    doc.save(str(base))

    job = _seed(session)
    profile = session.query(CandidateProfile).one()
    profile.raw_resume_path = str(base)
    tailored = TailoredResume(job_id=job.id, summary_text="Engineer who ships.", skills_order=["Java"],
                              validation_status="PASSED", docx_path="/somewhere/else/tailored_resume.docx")
    session.add(tailored)
    session.flush()
    session.add(TailoredBulletProvenance(tailored_resume_id=tailored.id, generated_text="Did the thing, twice as fast",
                                         source_bullet_id=profile.employments[0].bullets[0].id, validation_passed=True))
    session.commit()

    path = rebuild_tailored_resume(session, tailored, AppConfig(candidate=CandidateConfig(name="J", base_resume="x")))

    assert path is not None and path.parent.name.endswith(f"_{job.id}")  # the job's own folder, not the stale path
    out = Document(str(path))
    assert [p.text for p in out.paragraphs] == ["Engineer who ships.", "Did the thing, twice as fast"]
    assert out.paragraphs[1].style.name == "List Bullet"
    assert tailored.docx_path == str(path)
