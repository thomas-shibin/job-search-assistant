"""The Canadian resume check page end to end: the findings, converting a Word
resume in its own design (a new version), title equivalents (typed or
suggested), and a PDF resume moving into the standard layout."""
from __future__ import annotations

import io
import shutil
import struct
import zlib
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Inches
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

import app.models  # noqa: F401  register all models on Base.metadata
from app.core.database.base import Base
from app.core.database.session import PROJECT_ROOT, get_session
from app.modules.resumes import canadian_convert
from app.modules.resumes.models import CandidateProfile
from app.modules.resumes.parser import import_resume
from app.modules.usage.ai_calls import AICallLog
from app.modules.usage.models import ActivityEvent, ActivityKind

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"


def _png(size: int = 120) -> bytes:
    raw = b"".join(b"\x00" + b"\x88\x66\x44" * size for _ in range(size))

    def chunk(kind: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)

    header = struct.pack(">IIBBBBB", size, size, 8, 2, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b"")


def build_newcomer_resume(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Priya Sharma")
    doc.add_picture(io.BytesIO(_png()), width=Inches(1.2))
    doc.add_paragraph().add_run().add_picture(io.BytesIO(_png(16)), width=Inches(0.2))  # a contact icon, not a photo
    doc.add_paragraph("priya.sharma@example.com | +1 416 555 0199 | Toronto, ON")
    doc.add_heading("Summary", 1)
    doc.add_paragraph("HR professional who organised hiring for a large organisation.")
    doc.add_heading("Experience", 1)
    title = doc.add_paragraph()
    title.add_run("Senior HR Executive").bold = True
    title.add_run(", Tata Motors | 03/2019 - 05/2023")
    doc.add_paragraph("Analysed labour turnover across 4 plants and travelled to each one", style="List Bullet")
    doc.add_heading("Education", 1)
    doc.add_paragraph("Bachelor of Commerce, Pune University, 2018")
    doc.add_heading("Personal Details", 1)
    doc.add_paragraph("Date of Birth: 12/05/1995")
    doc.add_paragraph("Marital Status: Married")
    doc.add_paragraph("Languages: English, Hindi, French")
    doc.add_heading("Declaration", 1)
    doc.add_paragraph("I hereby declare that the above information is true to the best of my knowledge.")
    doc.add_paragraph("Place: Pune")
    doc.save(str(path))


@pytest.fixture()
def world(tmp_path):
    engine = create_engine(f"sqlite:///{tmp_path / 'test.db'}")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False, autocommit=False)

    def override_get_session():
        session = factory()
        try:
            yield session
        finally:
            session.close()

    import main

    main.app.dependency_overrides[get_session] = override_get_session
    yield TestClient(main.app), factory, tmp_path
    main.app.dependency_overrides.pop(get_session, None)


def _import(factory, path: Path) -> None:
    with factory() as s:
        import_resume(s, path)


@pytest.fixture()
def newcomer(world):
    http, factory, tmp_path = world
    resume = tmp_path / "priya.docx"
    build_newcomer_resume(resume)
    _import(factory, resume)
    return http, factory


def _active(factory) -> CandidateProfile:
    with factory() as s:
        profile = s.query(CandidateProfile).filter_by(is_active=True).one()
        _ = [(e.title, e.bullets) for e in profile.employments], profile.education
        s.expunge_all()
        return profile


def _document(profile: CandidateProfile) -> Document:
    path = Path(profile.raw_resume_path)
    return Document(str(path if path.is_absolute() else PROJECT_ROOT / path))


def _text(document: Document) -> str:
    return "\n".join(p.text for p in document.paragraphs)


def _pictures(document: Document) -> int:
    return len(document.inline_shapes)


def test_the_page_lists_what_to_change(newcomer):
    http, _ = newcomer
    page = http.get("/resume/canadian")
    assert page.status_code == 200
    text = page.text
    for title in ("Remove the photo", "Remove personal details", "Remove the declaration",
                  "Write dates as month and year", "Use Canadian spelling"):
        assert title in text
    assert "<del>Date of Birth: 12/05/1995</del>" in text
    # Each change is shown in its line, under its heading, with its own tick box.
    assert "<del>03/2019</del> <ins>Mar 2019</ins>" in text and "<ins>organization.</ins>" in text
    assert 'name="change" value="photo:-1"' in text and 'name="change" value="dates:' in text
    assert 'data-section="Experience"' in text and 'data-section="Personal Details"' in text
    assert "Senior HR Executive" in text and 'name="title-' in text
    assert "Suggest equivalents" not in text  # the free mock provider can't suggest titles


def test_converting_keeps_the_design_and_saves_a_new_version(newcomer):
    http, factory = newcomer
    before = _active(factory)
    role_id = before.employments[0].id
    data = {"fix": ["photo", "personal", "declaration", "dates", "spelling"], f"title-{role_id}": "HR Generalist"}

    response = http.post("/resume/canadian", data=data, follow_redirects=False)

    assert response.status_code == 303 and "converted=1" in response.headers["location"]
    after = _active(factory)
    assert after.version == before.version + 1
    assert after.employments[0].title == "Senior HR Executive (HR Generalist)"
    document = _document(after)
    text = _text(document)
    assert "Senior HR Executive (HR Generalist), Tata Motors | Mar 2019 - May 2023" in text
    assert "organized hiring for a large organization" in text
    assert "Analyzed labour turnover across 4 plants and travelled to each one" in text
    assert "Languages: English, Hindi, French" in text and "Personal Details" in text  # something useful stays
    for gone in ("Date of Birth", "Marital Status", "I hereby declare", "Place: Pune", "Declaration"):
        assert gone not in text
    assert _pictures(document) == 1  # the photo went; the small icon stayed
    title = next(p for p in document.paragraphs if p.text.startswith("Senior HR Executive"))
    assert title.runs[0].bold and title.runs[0].text.startswith("Senior HR Executive")  # formatting kept
    with factory() as s:
        events = s.query(ActivityEvent).filter_by(kind=ActivityKind.RESUME_CANADIAN).all()
        assert len(events) == 1 and "removed the photo" in events[0].summary

    page = http.get(response.headers["location"])
    assert "Your converted resume is ready" in page.text and "Remove the photo" not in page.text


def test_only_the_chosen_fixes_are_made(newcomer):
    http, factory = newcomer
    http.post("/resume/canadian", data={"fix": ["photo"]})
    document = _document(_active(factory))
    assert _pictures(document) == 1  # photo removed
    assert "Date of Birth: 12/05/1995" in _text(document) and "03/2019" in _text(document)


def test_nothing_chosen_changes_nothing(newcomer):
    http, factory = newcomer
    before = _active(factory)
    response = http.post("/resume/canadian", data={}, follow_redirects=False)
    assert response.status_code == 303 and "error=" in response.headers["location"]
    assert _active(factory).version == before.version


class _Provider:
    def __init__(self, answer):
        self.answer, self.last_usage = answer, None

    def canadian_titles(self, roles):
        self.roles = roles
        return {role["id"]: self.answer for role in roles}


def test_suggested_titles_fill_the_boxes_and_count_as_an_ai_call(newcomer, monkeypatch):
    http, factory = newcomer
    monkeypatch.setattr(canadian_convert, "build_ai_provider", lambda config, settings: _Provider("HR Generalist"))
    role_id = _active(factory).employments[0].id

    page = http.post("/resume/canadian/titles", data={})

    assert page.status_code == 200
    assert f'name="title-{role_id}" maxlength="60" value="HR Generalist"' in page.text
    assert "Suggested a Canadian equivalent for 1 job title" in page.text
    with factory() as s:
        assert s.query(AICallLog).filter_by(operation="canadian_titles").count() == 1
        assert s.query(ActivityEvent).filter_by(kind=ActivityKind.CANADIAN_TITLES).count() == 1


def test_more_senior_suggestions_are_dropped_and_typed_titles_kept(newcomer, monkeypatch):
    http, factory = newcomer
    monkeypatch.setattr(canadian_convert, "build_ai_provider", lambda config, settings: _Provider("HR Director"))
    role_id = _active(factory).employments[0].id

    page = http.post("/resume/canadian/titles", data={})
    assert "already clear to Canadian employers" in page.text and "HR Director" not in page.text

    typed = http.post("/resume/canadian/titles", data={f"title-{role_id}": "People Coordinator"})
    assert 'value="People Coordinator"' in typed.text


def test_a_pdf_resume_moves_into_the_standard_layout(world):
    http, factory, tmp_path = world
    resume = tmp_path / "resume.pdf"
    shutil.copy(FIXTURES / "sample_resume.pdf", resume)
    _import(factory, resume)
    before = _active(factory)

    page = http.get("/resume/canadian")
    assert page.status_code == 200 and "standard layout" in page.text

    response = http.post("/resume/canadian", data={"fix": ["spelling"]}, follow_redirects=False)
    assert response.status_code == 303 and "converted=1" in response.headers["location"]
    after = _active(factory)
    assert after.version == before.version + 1 and after.raw_resume_path.endswith(".docx")
    assert [e.title for e in after.employments] == [e.title for e in before.employments]


def test_no_resume_yet_goes_to_upload(world):
    http, _, _ = world
    response = http.get("/resume/canadian", follow_redirects=False)
    assert response.status_code == 303 and response.headers["location"].startswith("/resume?")


def test_tables_are_fixed_for_ats_software_when_chosen(world):
    http, factory, tmp_path = world
    resume = tmp_path / "tables.docx"
    doc = Document()
    doc.add_paragraph("Jordan Lee")
    doc.add_paragraph("jordan.lee@example.com | 416 555 0199 | Toronto, ON")
    doc.add_heading("Experience", 1)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Financial Analyst, Acme Ltd", "Jan 2020 – Present"
    table.cell(1, 0).text = "Rebuilt monthly reporting for five regional teams, cutting the close from ten days to six"
    doc.add_heading("Education", 1)
    doc.add_paragraph("Bachelor of Commerce, University of Toronto, 2018")
    doc.save(str(resume))
    _import(factory, resume)
    before = _active(factory)

    page = http.get("/resume/canadian").text
    assert "Move your content out of tables" in page and 'name="change" value="ats_tables:0"' in page
    assert "Financial Analyst, Acme Ltd | Jan 2020 – Present" in page  # the table now, then as plain lines

    response = http.post("/resume/canadian", data={"change": ["ats_tables:0"]}, follow_redirects=False)

    assert response.status_code == 303 and "converted=1" in response.headers["location"]
    after = _active(factory)
    assert after.version == before.version + 1
    document = _document(after)
    assert not document.tables
    assert "Financial Analyst, Acme Ltd\tJan 2020 – Present" in _text(document)
    page = http.get(response.headers["location"]).text
    assert "Turned 1 table into plain lines" in page and "Move your content out of tables" not in page


def test_single_changes_can_be_picked_and_the_rest_stay(newcomer):
    http, factory = newcomer
    page = http.get("/resume/canadian").text
    import re as _re
    dates = _re.findall(r'name="change" value="(dates:\d+)"', page)
    spelling = _re.findall(r'name="change" value="(spelling:\d+)"', page)
    assert dates and len(spelling) >= 2

    # One spelling change, nothing else.
    response = http.post("/resume/canadian", data={"change": [spelling[0]]}, follow_redirects=False)

    assert response.status_code == 303 and "converted=1" in response.headers["location"]
    text = _text(_document(_active(factory)))
    assert "03/2019" in text  # the dates weren't chosen
    assert ("organized hiring" in text) != ("Analyzed labour" in text)  # exactly one of the two spelling lines
    assert _pictures(_document(_active(factory))) == 2  # nor the photo


def test_one_table_of_two_can_be_picked(world):
    http, factory, tmp_path = world
    resume = tmp_path / "two_tables.docx"
    doc = Document()
    doc.add_paragraph("Jordan Lee")
    doc.add_paragraph("jordan.lee@example.com | 416 555 0199 | Toronto, ON")
    doc.add_heading("Experience", 1)
    for title in ("Financial Analyst, Acme Ltd", "Junior Analyst, Bolt Inc"):
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = title, "2019 – 2022"
        table.cell(1, 0).text = "Rebuilt monthly reporting for five regional teams, cutting the close from ten to six days"
    doc.save(str(resume))
    _import(factory, resume)

    http.post("/resume/canadian", data={"change": ["ats_tables:1"]})

    document = _document(_active(factory))
    assert len(document.tables) == 1 and document.tables[0].cell(0, 0).text == "Financial Analyst, Acme Ltd"
    assert "Junior Analyst, Bolt Inc\t2019 – 2022" in _text(document)


def test_compare_versions_and_go_back_any_time(newcomer):
    http, factory = newcomer
    first = _active(factory)
    http.post("/resume/canadian", data={"fix": ["declaration", "spelling"]})
    second = _active(factory)
    assert second.version == first.version + 1

    after = http.get("/resume/canadian?converted=1").text
    assert f"/resume/compare?from={first.version}&to={second.version}" in after
    assert f'action="/resume/versions/{first.version}/use"' in after  # undo

    compare = http.get(f"/resume/compare?from={first.version}&to={second.version}")
    assert compare.status_code == 200
    assert "<del>organised</del>" in compare.text and "<ins>organized</ins>" in compare.text
    assert "Removed" in compare.text and "I hereby declare" in compare.text
    assert "unchanged line" in compare.text  # only the changes, by default
    assert "unchanged line" not in http.get(f"/resume/compare?from={first.version}&to={second.version}&all=1").text

    download = http.get(f"/resume/versions/{first.version}/file")
    assert download.status_code == 200 and download.content[:2] == b"PK"

    back = http.post(f"/resume/versions/{first.version}/use", data={"next": "/resume"}, follow_redirects=False)
    assert back.status_code == 303 and "in+use+again" in back.headers["location"]
    assert _active(factory).id == first.id
    with factory() as s:
        assert s.query(CandidateProfile).count() == 2  # nothing deleted

    # A change after going back is numbered after the newest version, never a second "version 2".
    http.post("/resume/canadian", data={"fix": ["photo"]})
    assert _active(factory).version == second.version + 1
    history = http.get("/resume").text
    assert "Use this version" in history and f"/resume/compare?to={second.version + 1}" in history


def test_a_version_that_isnt_yours_or_doesnt_exist_is_refused(newcomer):
    http, _ = newcomer
    assert http.get("/resume/versions/99/file").status_code == 404
    gone = http.post("/resume/versions/99/use", data={"next": "https://evil.example"}, follow_redirects=False)
    assert gone.headers["location"].startswith("/resume?error=")  # and never off this site
