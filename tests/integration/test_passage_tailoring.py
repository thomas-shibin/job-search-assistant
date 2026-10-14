"""Word resumes are tailored paragraph by paragraph whatever their layout --
including project-style resumes whose work is described in tables."""
from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

import app.models  # noqa: F401  register all models on Base.metadata
from app.integrations.ai.mock import MockAIProvider
from app.core.config.schema import AIConfig, AppConfig, CandidateConfig
from app.core.database.base import Base
from app.documents.docx_edit import is_prose, rewritable_passages
from app.ingestion.pipeline import ingest_raw_job
from app.models.candidate import CandidateProfile
from app.resume.tailor import tailor_resume
from app.resume.validator import validate_passage
from app.services.jobs import get_job_detail
from app.integrations.job_boards.base import RawJob

SUMMARY = "Full-stack developer with 5 years of experience delivering web apps in React and Node."
DUTY = "Built the reporting dashboards in React that operations teams use every day"
BACKEND_DUTY = "Wrote the Node services behind the dashboards and kept their APIs fast and reliable"


def _project_resume(path: Path, with_backend_duty: bool = False) -> None:
    doc = Document()
    doc.add_paragraph("Aidan Lazar")
    doc.add_paragraph("aidan@example.com | (250) 555-0100")
    doc.add_paragraph("PROFESSIONAL SUMMARY")
    doc.add_paragraph(SUMMARY)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "PROJECT NAME", "Salt Security"
    table.cell(1, 0).text, table.cell(1, 1).text = "DURATION", "March 2024 – March 2026"
    doc.add_paragraph("Responsibilities")
    doc.add_paragraph(DUTY, style="List Bullet")
    if with_backend_duty:
        doc.add_paragraph(BACKEND_DUTY, style="List Bullet")
    doc.add_paragraph("Languages: Java | Go | Python | React | Node")
    doc.save(str(path))


@pytest.fixture()
def session(tmp_path, monkeypatch):
    monkeypatch.setenv("APPLICATIONS_DATA_DIR", str(tmp_path / "applications"))
    engine = create_engine(f"sqlite:///{tmp_path / 't.db'}")
    Base.metadata.create_all(engine)
    with Session(engine) as s:
        yield s


def test_only_prose_is_rewritable():
    assert is_prose(SUMMARY) and is_prose(DUTY)
    assert is_prose("Cloud Computing: Skilled in deploying web applications to AWS using GitHub Actions and Docker.")
    for line in ("Senior Application Developer, 05/2021 – Present", "Bachelor of Engineering — Computer Science",
                 "AWS Certified Solutions Architect — Associate", "Languages: Java | Go | Python | Spring Boot",
                 "+1 (236) 334-7863 | me@example.com | Victoria, Canada", "Technical Lead and FullStack Application Developer"):
        assert not is_prose(line), line


def test_rewrites_that_claim_something_new_keep_the_original():
    resume = f"{SUMMARY}\n{DUTY}"
    assert validate_passage(DUTY, "Built React reporting dashboards used daily by operations", resume_text=resume)[1] is None
    for rewrite in ("Built React dashboards used by 40 operations teams",           # new number
                    "Built React and Angular dashboards for operations teams",      # new tool
                    "As a Staff Engineer, built React dashboards for operations"):  # new title
        text, reason = validate_passage(DUTY, rewrite, resume_text=resume, job_title="Staff Engineer")
        assert text == DUTY and reason


def test_a_project_style_resume_is_tailored_in_place(session, tmp_path):
    base = tmp_path / "resume.docx"
    _project_resume(base)
    assert [p.text for p in rewritable_passages(base)] == [SUMMARY, DUTY]
    session.add(CandidateProfile(raw_resume_path=str(base), source_hash="h", summary=None, is_active=True))
    job = ingest_raw_job(session, RawJob(source="manual", source_url="https://x/1", company="Acme",
                                         title="Frontend Developer", description="React dashboards. " * 20)).job
    session.commit()

    class _Provider(MockAIProvider):
        def rewrite_passages(self, passages, job, requirements, candidate_skills):
            return {p["id"]: ("React full-stack developer with 5 years delivering web apps with Node."
                              if p["section"] == "PROFESSIONAL SUMMARY" else
                              "Built React reporting dashboards that operations teams rely on daily") for p in passages}

    result = tailor_resume(session, job, AppConfig(candidate=CandidateConfig(name="", base_resume=str(base)),
                                                   ai=AIConfig(provider="mock")), ai_provider=_Provider())

    before, after = Document(str(base)), Document(str(result.docx_path))
    assert len(after.paragraphs) == len(before.paragraphs)
    texts = [p.text for p in after.paragraphs]
    assert "React full-stack developer with 5 years delivering web apps with Node." in texts
    assert "Built React reporting dashboards that operations teams rely on daily" in texts
    unchanged = ("Aidan Lazar", "aidan@example.com | (250) 555-0100", "PROFESSIONAL SUMMARY", "Responsibilities",
                 "Languages: Java | Go | Python | React | Node")
    assert all(line in texts for line in unchanged)
    assert after.tables[0].cell(1, 1).text == "March 2024 – March 2026"

    diffs = get_job_detail(session, job.id).bullet_diffs  # "What changed" lists every rewritten paragraph
    assert [(d.role, d.source_text) for d in diffs] == [("PROFESSIONAL SUMMARY", SUMMARY), ("Responsibilities", DUTY)]
    assert all(d.validation_passed for d in diffs)


def test_downloads_are_named_after_the_candidate_company_and_time(monkeypatch):
    import time
    from datetime import datetime
    from types import SimpleNamespace

    from app.resume.tailor import resume_download_name

    monkeypatch.setenv("TZ", "America/Vancouver")
    time.tzset()
    config = AppConfig(candidate=CandidateConfig(name="", base_resume="x.docx"))
    try:
        name = resume_download_name(config, SimpleNamespace(full_name="GLADIA JOY"), "Acme / Sons: Ltd",
                                    datetime(2026, 10, 10, 21, 32), ".docx")  # stored in UTC
    finally:
        monkeypatch.delenv("TZ")
        time.tzset()
    assert name == "Gladia Joy - Acme Sons Ltd - 2026-10-10_1432.docx"


def test_suggested_removals_only_happen_once_the_user_approves(session, tmp_path):
    from app.models.resume import TailoredBulletProvenance
    from app.resume.tailor import rebuild_tailored_resume

    base = tmp_path / "resume.docx"
    _project_resume(base, with_backend_duty=True)
    session.add(CandidateProfile(raw_resume_path=str(base), source_hash="h", summary=None, is_active=True))
    job = ingest_raw_job(session, RawJob(source="manual", source_url="https://x/2", company="Acme",
                                         title="Backend Developer", description="Java services. " * 20)).job
    session.commit()

    class _Provider(MockAIProvider):
        def rewrite_passages(self, passages, job, requirements, candidate_skills):
            return {p["id"]: ({"text": p["text"], "drop_reason": "Front-end work; this role is backend"}
                              if p["text"] in (DUTY, SUMMARY) else p["text"]) for p in passages}

    config = AppConfig(candidate=CandidateConfig(name="", base_resume=str(base)), ai=AIConfig(provider="mock"))
    result = tailor_resume(session, job, config, ai_provider=_Provider())

    # Suggested, not removed: the resume still has the paragraph.
    assert DUTY in [p.text for p in Document(str(result.docx_path)).paragraphs]
    diffs = get_job_detail(session, job.id).bullet_diffs
    assert not next(d for d in diffs if d.source_text == SUMMARY).suggest_remove  # never the summary
    diff = next(d for d in diffs if d.source_text == DUTY)
    assert diff.suggest_remove and diff.remove_reason == "Front-end work; this role is backend" and not diff.removed

    # The user approves: rebuilt from the base without it, everything else intact.
    row = session.get(TailoredBulletProvenance, diff.id)
    row.removed = True
    session.commit()
    rebuild_tailored_resume(session, result.tailored_resume, config)
    after = [p.text for p in Document(str(result.docx_path)).paragraphs]
    assert DUTY not in after and "Responsibilities" in after and SUMMARY in after
    assert len(after) == len(Document(str(base)).paragraphs) - 1

    # ...and changes their mind: it's back.
    row.removed = False
    session.commit()
    rebuild_tailored_resume(session, result.tailored_resume, config)
    assert DUTY in [p.text for p in Document(str(result.docx_path)).paragraphs]


def test_removing_the_only_paragraph_of_a_table_cell_keeps_the_table_valid(tmp_path):
    from app.documents.docx_edit import tailor_docx

    base, out = tmp_path / "base.docx", tmp_path / "out.docx"
    doc = Document()
    doc.add_table(rows=1, cols=1).cell(0, 0).text = DUTY
    doc.save(str(base))

    report = tailor_docx(base, out, [], remove=[DUTY])

    cell = Document(str(out)).tables[0].cell(0, 0)
    assert report.removed == 1 and cell.text == "" and len(cell.paragraphs) == 1


def test_a_whole_section_is_never_suggested_for_removal():
    from app.documents.docx_edit import Passage
    from app.resume.tailor import _sensible_suggestions

    a, b, c = Passage("x", "Project A"), Passage("y", "Project A"), Passage("z", "Project B")
    kept = _sensible_suggestions([(a, "x", None, "UI work"), (b, "y", None, "UI work"), (c, "z", None, "UI work")])
    assert [drop for *_, drop in kept] == [None, None, None]  # each would empty its section
    kept = _sensible_suggestions([(a, "x", None, "UI work"), (b, "y", None, None)])
    assert [drop for *_, drop in kept] == ["UI work", None]
