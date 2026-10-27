"""The free Canadian resume check anyone can use (/canadian-resume-check):
checking without an account, files that can't be checked, limits, and
carrying the checked resume into an account to fix it without uploading again."""
from __future__ import annotations

import html
import io
import zipfile
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

import app.models  # noqa: F401  register all models on Base.metadata
from app.core.database.base import Base
from app.core.database.session import get_session
from app.modules.auth.ratelimit import RateLimiter
from app.modules.resumes import public_check, uploads
from app.modules.resumes.models import CandidateProfile
from tests.integration.test_canadian_resume import build_newcomer_resume

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture()
def world(tmp_path, monkeypatch):
    engine = create_engine(f"sqlite:///{tmp_path / 'test.db'}")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False, autocommit=False)

    def override_get_session():
        session = factory()
        try:
            yield session
        finally:
            session.close()

    import main

    public_check.clear()
    public_check.limiter.reset()
    main.app.dependency_overrides[get_session] = override_get_session
    resume = tmp_path / "priya.docx"
    build_newcomer_resume(resume)
    yield TestClient(main.app), factory, resume.read_bytes()
    main.app.dependency_overrides.pop(get_session, None)
    public_check.clear()


def _check(http, data: bytes, name: str = "Priya resume.docx", kind: str = DOCX):
    response = http.post("/canadian-resume-check", files={"resume": (name, data, kind)}, follow_redirects=False)
    response.page = html.unescape(response.text)
    return response


@pytest.mark.real_auth
def test_anyone_can_open_the_page_and_search_engines_may_index_it(world):
    http, _, _ = world
    page = http.get("/canadian-resume-check")
    assert page.status_code == 200
    assert "Free · No sign-up" in page.text and 'name="resume"' in page.text
    assert 'rel="canonical"' in page.text and "noindex" not in page.text
    # Three expectations, then a pointer to the check rather than the whole list.
    guide = page.text.split('class="check-guide-list"', 1)[1].split("</section>", 1)[0]
    assert guide.count("<li>") == 3 and "And 5 more things employers look for" in guide
    assert "<details" not in guide
    # While it checks, six steps worded indirectly: they don't name what it looks at.
    steps = page.text.split('<ol class="check-steps"', 1)[1].split("</ol>", 1)[0]
    assert "data-check-steps" in page.text and steps.count("<li>") == 6
    # Reached by a link, htmx would send the form itself and skip the wait.
    assert 'hx-boost="false" data-check-steps' in page.text
    assert not any(word in steps for word in ("Photo", "personal", "References", "ATS", "Dates", "spelling"))


@pytest.mark.real_auth
def test_a_visitor_gets_scores_but_not_what_to_change_without_an_account(world):
    http, _, data = world
    response = _check(http, data)
    text = response.text
    assert response.status_code == 200
    # How many problems and how serious, but not what they are: none of it is in the page at all.
    review = public_check.check("Priya resume.docx", data)
    assert text.count('class="locked-finding level-') == len(review.findings) + len(review.ats)
    assert f"We found {len(review.findings)} things to look at" in text and "See what to change" in text
    for f in review.findings + review.ats:
        assert f.title not in text and f.advice not in text
    assert "What happens:" not in text and "Date of Birth" not in text and "Your resume, marked up" not in text
    assert "/100" in text and "Needs work for Canadian employers" in text and "Priya resume.docx" in text
    assert 'name="fix"' not in text  # nothing to tick: fixing needs an account
    assert "Create a free account to fix it" in text and 'data-next="/resume/canadian"' in text
    assert public_check.COOKIE in response.headers["set-cookie"] and "HttpOnly" in response.headers["set-cookie"]
    assert response.headers["cache-control"] == "no-store"  # their resume's text is on the page
    assert not uploads.UPLOAD_DIR.exists() or not any(uploads.UPLOAD_DIR.rglob("*"))  # nothing saved to disk


def test_a_member_sees_each_problem_and_what_it_costs_without_a_reference_list(world):
    http, _, data = world
    page = _check(http, data).page
    assert "Remove the photo" in page and "Remove personal details" in page
    assert "What happens:" in page and "can lower your chances of getting the job" in page
    # Sources stay behind the checks (check_sources.py), not listed link by link for visitors.
    assert "jobbank.gc.ca" not in page and "Read the source" not in page
    findings = page.split("What Canadian employers check", 1)[1].split("Will ATS software", 1)[0]
    assert findings.count('<article class="finding') <= 5
    assert "data-check-results" in page  # hidden in the browser when another resume is checked


def test_the_five_most_important_problems_come_first_and_the_rest_behind_show_more():
    from app.modules.resumes.canadian import MUST, Finding
    from app.web.templating import templates

    macros = templates.env.get_template("canadian_check_public.html").module
    items = [Finding(f"k{n}", MUST, f"Problem {n}", "Advice") for n in range(7)]
    shown, _, more = str(macros.findings_list(items)).partition('<details class="check-more">')
    assert shown.count("<article") == 5 and "Problem 4" in shown and "Problem 5" not in shown
    assert "Show 2 more" in more and more.count("<article") == 2
    assert '<details class="check-more">' not in str(macros.findings_list(items[:5]))


@pytest.mark.real_auth
def test_a_visitor_sees_the_measured_score_after_the_fixes_but_not_the_fixed_text(world):
    http, _, data = world
    page = _check(http, data).page
    review = public_check.check("Priya resume.docx", data)
    assert review.after_score > review.score  # measured on the fixed text, not assumed
    assert not any(f.key in ("photo", "personal", "declaration") for f in review.after)
    assert f"{review.score} → {review.after_score}" in page
    assert "Your resume after the fixes" in page and "is-skeleton" in page and "Get my fixed resume" in page
    assert not any(line.text in page for line in review.preview if len(line.text) > 12)
    seconds = int(page.split('data-hold-seconds="')[1].split('"')[0])
    assert public_check.HOLD_SECONDS - 5 <= seconds <= public_check.HOLD_SECONDS


def test_a_member_sees_their_resume_marked_up_and_the_fixed_version(world):
    http, _, data = world
    page = _check(http, data).page
    review = public_check.check("Priya resume.docx", data)
    assert "Your resume, marked up" in page and "resume-line mark-must" in page
    assert "Date of Birth: 12/05/1995" in page  # their own line, marked
    assert all("Date of Birth" not in line.text for line in review.preview)
    assert all(line.text in page for line in review.preview)


@pytest.mark.real_auth
def test_the_ats_check_is_honest_about_what_it_can_tell(world):
    http, _, data = world
    page = _check(http, data).page
    assert "Will ATS software read it correctly?" in page and "No tool can promise" in page
    assert "/100" in page and "ATS readability" in page


def _ready_docx() -> bytes:
    doc, buffer = Document(), io.BytesIO()
    doc.add_paragraph("Jordan Lee")
    doc.add_paragraph("jordan.lee@example.com | 416 555 0199 | Toronto, ON")
    doc.add_heading("Summary", 1)
    doc.add_paragraph("Financial analyst with six years in retail banking, strong in forecasting and reporting.")
    doc.add_heading("Experience", 1)
    doc.add_paragraph("Financial Analyst, Maple Bank, Jan 2019 - Present")
    for _ in range(3):
        doc.add_paragraph("Rebuilt monthly reporting for five branches, cutting the close from ten days to six",
                          style="List Bullet")
    doc.add_heading("Education", 1)
    doc.add_paragraph("Bachelor of Commerce, University of Toronto, 2018")
    doc.save(buffer)
    return buffer.getvalue()


@pytest.mark.real_auth
def test_a_ready_resume_is_told_so_and_pointed_at_jobs(world):
    http, _, _ = world
    response = _check(http, _ready_docx(), "jordan.docx")
    assert "Your resume is ready for Canada" in response.page and "See my match score on every job" in response.page
    assert "Create a free account to fix it" not in response.page and "Your resume after the fixes" not in response.page
    assert public_check.COOKIE in response.headers["set-cookie"]  # kept, to match jobs after signing up


def test_a_member_with_a_ready_resume_goes_straight_to_jobs(world):
    http, factory, _ = world
    checked = _check(http, _ready_docx(), "jordan.docx")
    assert 'name="next" value="/jobs"' in checked.text
    claimed = http.post("/resume/canadian/claim", data={"next": "/jobs"}, follow_redirects=False)
    assert claimed.headers["location"].startswith("/jobs?msg=")
    _check(http, _ready_docx(), "jordan.docx")
    elsewhere = http.post("/resume/canadian/claim", data={"next": "https://evil.example"}, follow_redirects=False)
    assert elsewhere.headers["location"].startswith("/resume/canadian?msg=")  # only known pages


@pytest.mark.real_auth
def test_a_pdf_is_checked_too(world):
    http, _, _ = world
    response = _check(http, (FIXTURES / "sample_resume.pdf").read_bytes(), "cv.pdf", "application/pdf")
    assert response.status_code == 200 and "/100" in response.text and "cv.pdf" in response.text


def _zip(names: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, body in names.items():
            archive.writestr(name, body)
    return buffer.getvalue()


def _short_docx() -> bytes:
    doc, buffer = Document(), io.BytesIO()
    doc.add_paragraph("Priya Sharma")
    doc.save(buffer)
    return buffer.getvalue()


@pytest.mark.real_auth
@pytest.mark.parametrize("name, data, message", [
    ("resume.docx", b"", "That file is empty"),
    ("notes.txt", b"just some text", "Choose your resume as a PDF or Word (.docx) file"),
    ("old.doc", b"\xd0\xcf\x11\xe0" + b"\0" * 100, "old Word (.doc) file"),
    ("photos.docx", None, "isn't a Word document"),
    ("broken.docx", b"PK\x03\x04 not really a zip", "damaged"),
    ("broken.pdf", b"%PDF-1.4 nothing else", "Couldn't read that file"),
    ("short.docx", None, "almost no text"),
])
def test_files_that_cant_be_checked_are_explained(world, name, data, message):
    http, _, _ = world
    data = {"photos.docx": _zip({"photo.jpg": b"x"}), "short.docx": _short_docx()}.get(name, data)
    response = _check(http, data, name)
    assert response.status_code == 400 and message in response.page
    assert "set-cookie" not in response.headers


@pytest.mark.real_auth
def test_size_and_page_limits(world, monkeypatch):
    http, _, data = world
    monkeypatch.setattr(public_check, "MAX_RESUME_BYTES", 1000)
    assert "over 0 MB" in _check(http, data).page
    monkeypatch.setattr(public_check, "MAX_RESUME_BYTES", 10_000_000)
    monkeypatch.setattr(public_check, "MAX_PAGES", 0)
    response = _check(http, (FIXTURES / "sample_resume.pdf").read_bytes(), "cv.pdf", "application/pdf")
    assert response.status_code == 400 and "should be one or two" in response.page
    monkeypatch.setattr(public_check, "_MAX_UNZIPPED", 10)
    assert "too large to check" in _check(http, data).page


@pytest.mark.real_auth
def test_checks_are_limited_per_visitor(world, monkeypatch):
    http, _, data = world
    monkeypatch.setattr(public_check, "limiter", RateLimiter(limit=1, window_seconds=3600))
    assert _check(http, data).status_code == 200
    assert _check(http, data).status_code == 429


@pytest.mark.real_auth
def test_switched_off_by_an_admin(world, monkeypatch):
    from app.modules.access import rbac

    http, _, data = world
    monkeypatch.setattr(rbac, "feature_on", lambda session, key: key != "canadian_resume")
    assert http.get("/canadian-resume-check", follow_redirects=False).status_code == 303
    assert _check(http, data).status_code == 303
    landing = http.get("/").text  # the landing page offers an account instead of an upload that would fail
    assert 'action="/canadian-resume-check"' not in landing and "Create a free account" in landing


def test_a_member_fixes_the_checked_resume_without_uploading_again(world):
    http, factory, data = world
    checked = _check(http, data)
    assert 'action="/resume/canadian/claim"' in checked.text and "Fix my resume" in checked.text
    page = http.get("/resume/canadian")  # no resume of their own yet: offered the checked one
    assert page.status_code == 200 and "Use the resume you just checked?" in page.text
    assert 'data-hold-seconds="' in page.text
    claimed = http.post("/resume/canadian/claim", follow_redirects=False)
    assert claimed.status_code == 303 and claimed.headers["location"].startswith("/resume/canadian?msg=")
    assert public_check.COOKIE in claimed.headers["set-cookie"]  # cleared
    with factory() as s:
        assert s.query(CandidateProfile).filter_by(is_active=True).count() == 1
    fixing = http.get("/resume/canadian")
    assert "Remove the photo" in fixing.text and 'name="change" value="photo:-1"' in fixing.text
    assert "Use the resume you just checked?" not in fixing.text
    assert "Will ATS software read it correctly?" in fixing.text and "What happens:" in fixing.text
    again = http.post("/resume/canadian/claim", follow_redirects=False)
    assert again.headers["location"].startswith("/resume?error=") and "no+longer+held" in again.headers["location"]


def test_without_a_checked_resume_the_member_page_asks_for_an_upload(world):
    http, _, _ = world
    page = http.get("/resume/canadian", follow_redirects=False)
    assert page.status_code == 303 and page.headers["location"].startswith("/resume?error=")


def test_held_files_expire_and_memory_is_bounded(monkeypatch):
    public_check.clear()
    clock = [1000.0]
    monkeypatch.setattr(public_check.time, "monotonic", lambda: clock[0])
    token = public_check.hold("a.docx", b"x")
    assert public_check.held_name(token) == "a.docx"
    clock[0] += public_check.HOLD_SECONDS + 1
    assert public_check.held_name(token) is None and public_check.take(token) is None
    monkeypatch.setattr(public_check, "_MAX_HELD", 2)
    first = public_check.hold("1.docx", b"1")
    clock[0] += 1
    second = public_check.hold("2.docx", b"2")
    clock[0] += 1
    public_check.hold("3.docx", b"3", replaces=None)
    assert public_check.held_name(first) is None and public_check.held_name(second) == "2.docx"
    replaced = public_check.hold("4.docx", b"4", replaces=second)
    assert public_check.held_name(second) is None and public_check.take(replaced) == ("4.docx", b"4")
    assert public_check.take("made-up") is None and public_check.take(None) is None
    timed = public_check.hold("5.docx", b"5")
    clock[0] += 600
    assert public_check.seconds_left(timed) == public_check.HOLD_SECONDS - 600
    assert public_check.seconds_left("made-up") == 0 and public_check.seconds_left(None) == 0
    public_check.clear()


@pytest.mark.real_auth
def test_the_data_deletion_page_explains_how_for_anyone(world):
    http, _, _ = world
    page = http.get("/data-deletion")
    assert page.status_code == 200
    text = html.unescape(page.text)
    assert "Delete my account" in text and "Apps and websites" in text and 'href="/account"' in text
    assert 'href="/data-deletion"' in http.get("/privacy").text


def test_a_member_can_say_no_thanks_and_the_checked_file_is_deleted(world):
    http, factory, data = world
    _check(http, data)
    assert "Use the resume you just checked?" in http.get("/resume/canadian").text

    response = http.post("/resume/canadian/discard", follow_redirects=False)

    assert response.status_code == 303 and "Deleted+the+resume+you+checked" in response.headers["location"]
    assert public_check.COOKIE in response.headers["set-cookie"]  # cleared
    assert public_check.held_name(http.cookies.get(public_check.COOKIE)) is None
    with factory() as s:
        assert s.query(CandidateProfile).count() == 0  # nothing was imported


def test_checking_the_resume_already_in_use_offers_nothing(world):
    http, factory, data = world
    _check(http, data)
    http.post("/resume/canadian/claim")  # now it's their resume
    _check(http, data)  # and they check the same file again

    page = http.get("/resume/canadian").text

    assert "Use the resume you just checked?" not in page


def _health_card() -> bytes:
    from tests.unit.test_not_a_resume import NOT_RESUMES

    doc = Document()
    for line in NOT_RESUMES["health_card"].splitlines():
        doc.add_paragraph(line)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_a_file_that_isnt_a_resume_gets_no_score(world):
    http, _, _ = world
    response = _check(http, _health_card(), name="health card.docx")
    assert response.status_code == 400
    assert "doesn't look like a resume" in response.page
    assert "Canada-ready" not in response.page.split("doesn't look like a resume", 1)[1]
    assert public_check.COOKIE not in response.headers.get("set-cookie", "")  # nothing held


def test_a_member_cant_make_a_health_card_their_resume(world):
    http, factory, _ = world
    response = http.post("/resume", files={"resume": ("health card.docx", _health_card(), DOCX)},
                         follow_redirects=False)
    assert response.status_code == 303 and "look+like+a+resume" in response.headers["location"]
    with factory() as s:
        assert s.query(CandidateProfile).count() == 0
    assert not list(uploads.UPLOAD_DIR.rglob("*health_card*"))  # the file isn't kept


def test_an_older_non_resume_version_isnt_checked(world, tmp_path):
    http, factory, _ = world
    path = tmp_path / "card.docx"
    path.write_bytes(_health_card())
    from app.modules.resumes.parser import import_resume

    with factory() as s:  # as if uploaded before the check existed
        import_resume(s, path)
    page = http.get("/resume/canadian").text
    assert "doesn't look like a resume, so it wasn't checked" in page
    assert "Make the chosen changes" not in page and "things to change" not in page
