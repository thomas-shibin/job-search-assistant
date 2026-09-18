"""Generates tests/fixtures/sample_resume.docx.

Run directly to (re)build the fixture: python tests/fixtures/build_sample_resume.py
The fixture is checked in as a binary file so tests don't depend on regenerating it,
but this script is kept so its contents are transparent and reproducible.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document

OUTPUT_PATH = Path(__file__).resolve().parent / "sample_resume.docx"


def build(path: Path) -> None:
    doc = Document()

    doc.add_paragraph("Jordan Casey")
    doc.add_paragraph("jordan.casey@example.com | (555) 123-4567 | Victoria, BC")

    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph(
        "Senior backend engineer with 10+ years of experience building distributed "
        "systems in Java and Python, specializing in scalable REST APIs and "
        "cloud-native architectures."
    )

    doc.add_heading("Experience", level=1)

    _bold_paragraph(doc, "Senior Software Engineer, Acme Cloud Systems | Jan 2021 - Present")
    doc.add_paragraph(
        "Led design and implementation of a microservices platform handling 50M+ requests/day",
        style="List Bullet",
    )
    doc.add_paragraph(
        "Reduced p95 API latency by 40% through caching and query optimization",
        style="List Bullet",
    )
    doc.add_paragraph(
        "Mentored a team of 5 engineers and led architecture review sessions",
        style="List Bullet",
    )

    _bold_paragraph(doc, "Software Engineer, Beta Systems Inc. | Jun 2018 - Dec 2020")
    doc.add_paragraph(
        "Built and maintained REST APIs serving the company's core billing platform",
        style="List Bullet",
    )
    doc.add_paragraph(
        "Migrated a monolithic Java application to a set of containerized Spring Boot services",
        style="List Bullet",
    )
    doc.add_paragraph(
        "Introduced automated integration tests, cutting regression bugs by half",
        style="List Bullet",
    )

    _bold_paragraph(doc, "Junior Developer — Gamma Software (2015-2018)")
    doc.add_paragraph("Developed internal tooling for the QA team using Python", style="List Bullet")
    doc.add_paragraph("Fixed production defects reported through the support queue", style="List Bullet")
    doc.add_paragraph("Wrote unit tests for legacy modules with no prior coverage", style="List Bullet")

    doc.add_heading("Skills", level=1)
    doc.add_paragraph(
        "Java, Python, Spring Boot, REST APIs, Distributed Systems, AWS, Kubernetes, "
        "Docker, PostgreSQL, MongoDB, Kafka, OAuth2, Keycloak, Git, CI/CD, "
        "Microservices, Terraform, Linux"
    )

    doc.add_heading("Education", level=1)
    doc.add_paragraph("Bachelor of Science in Computer Science, University of Victoria (2011 - 2015)")

    doc.add_heading("Certifications", level=1)
    doc.add_paragraph("AWS Certified Solutions Architect – Associate, Amazon Web Services, 2022")
    doc.add_paragraph("Certified Kubernetes Administrator (CKA), Cloud Native Computing Foundation, 2023")

    doc.add_heading("Volunteer Experience", level=1)
    doc.add_paragraph(
        "Mentor, Local Coding Bootcamp (2019-Present) - helped junior developers with career transitions."
    )

    doc.save(str(path))


def _bold_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True


if __name__ == "__main__":
    build(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}")
