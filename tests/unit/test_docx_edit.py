"""In-place resume tailoring keeps the original document's formatting and
changes only the wording it is asked to."""
from __future__ import annotations

from docx import Document
from docx.shared import Pt, RGBColor

from app.modules.resumes.documents.docx_edit import set_paragraph_text, tailor_docx


def _base_resume(path):
    doc = Document()
    doc.sections[0].left_margin = Pt(36)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).paragraphs[0].add_run("Victoria, Canada")
    table.cell(0, 1).paragraphs[0].add_run("me@example.com")
    heading = doc.add_paragraph()
    run = heading.add_run("Summary")
    run.bold, run.font.size = True, Pt(13)
    summary = doc.add_paragraph()
    summary.add_run("Backend engineer building ")
    bold = summary.add_run("Kafka")
    bold.bold = True
    summary.add_run(" pipelines at scale.")
    doc.add_paragraph("Second summary paragraph.")
    bullet = doc.add_paragraph(style="List Bullet")
    styled = bullet.add_run("Built Java services for payments.")
    styled.font.size, styled.font.color.rgb, styled.font.name = Pt(12), RGBColor(0x11, 0x22, 0x33), "Georgia"
    doc.add_paragraph("• Typed-bullet line about Docker.")
    doc.save(str(path))


def _paragraphs(path):
    return [p for p in Document(str(path)).paragraphs if p.text.strip()]


def test_only_the_wording_changes(tmp_path):
    base, out = tmp_path / "base.docx", tmp_path / "out.docx"
    _base_resume(base)

    report = tailor_docx(base, out, [
        ("Built Java services for payments.", "Built high-throughput Java services for payments."),
        ("Typed-bullet line about Docker.", "Typed-bullet line about Docker and Kubernetes."),
    ])

    assert report.replaced == 2 and not report.not_found
    doc = Document(str(out))
    bullet = next(p for p in doc.paragraphs if "high-throughput" in p.text)
    assert bullet.style.name == "List Bullet"
    run = bullet.runs[0]
    assert (run.font.size, run.font.color.rgb, run.font.name) == (Pt(12), RGBColor(0x11, 0x22, 0x33), "Georgia")
    assert any(p.text == "• Typed-bullet line about Docker and Kubernetes." for p in doc.paragraphs)
    assert doc.sections[0].left_margin == Pt(36)
    assert len(doc.tables) == 1 and doc.tables[0].cell(0, 1).text == "me@example.com"
    heading = next(p for p in doc.paragraphs if p.text == "Summary")
    assert heading.runs[0].bold and heading.runs[0].font.size == Pt(13)


def test_bold_words_that_survive_a_rewrite_stay_bold(tmp_path):
    base = tmp_path / "base.docx"
    _base_resume(base)
    paragraph = next(p for p in Document(str(base)).paragraphs if "Kafka" in p.text)

    set_paragraph_text(paragraph, "Senior backend engineer building Kafka and Flink pipelines at scale.")

    assert paragraph.text == "Senior backend engineer building Kafka and Flink pipelines at scale."
    bold_words = " ".join(r.text for r in paragraph.runs if r.bold).split()
    assert "Kafka" in bold_words and "Senior" not in bold_words


def test_multi_paragraph_text_is_split_back_across_the_same_paragraphs(tmp_path):
    base, out = tmp_path / "base.docx", tmp_path / "out.docx"
    _base_resume(base)
    original = "Backend engineer building Kafka pipelines at scale. Second summary paragraph."
    before = len(_paragraphs(base))

    tailor_docx(base, out, [(original, "Backend engineer building Kafka pipelines for fintech. Mentors the team.")])

    texts = [p.text for p in _paragraphs(out)]
    assert len(texts) == before  # no paragraph removed
    assert "Backend engineer building Kafka pipelines for fintech." in texts
    assert "Mentors the team." in texts


def test_a_rewrite_that_cannot_fill_every_paragraph_is_left_as_written(tmp_path):
    base, out = tmp_path / "base.docx", tmp_path / "out.docx"
    _base_resume(base)
    original = "Backend engineer building Kafka pipelines at scale. Second summary paragraph."

    report = tailor_docx(base, out, [(original, "One merged sentence for both paragraphs.")])

    texts = [p.text for p in _paragraphs(out)]
    assert "Second summary paragraph." in texts and report.not_found == [original]


def test_location_line_keeps_its_shape_and_missing_text_is_reported(tmp_path):
    base, out = tmp_path / "base.docx", tmp_path / "out.docx"
    _base_resume(base)

    report = tailor_docx(base, out, [("A bullet that isn't there.", "Whatever.")], location_line="Toronto, ON, Canada")

    assert report.address_replaced and report.not_found == ["A bullet that isn't there."]
    assert Document(str(out)).tables[0].cell(0, 0).text == "Toronto, Canada"
