"""The ATS readability check (app/modules/resumes/ats.py), and that every check
the pages show has a basis (app/modules/resumes/check_sources.py)."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas

from app.modules.resumes import ats, canadian, check_sources

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"
BULLET = "Led a team of five analysts to rebuild monthly reporting, cutting the close from ten days to six"


def _keys(path: Path) -> set[str]:
    doc = canadian.read(path)
    return {f.key for f in ats.check(path, doc, path.stat().st_size)}


def _plain(doc: Document, contact: bool = True) -> Document:
    doc.add_paragraph("Jordan Lee")
    if contact:
        doc.add_paragraph("jordan.lee@example.com | 416 555 0199 | Toronto, ON")
    doc.add_heading("Experience", 1)
    for _ in range(4):
        doc.add_paragraph(BULLET, style="List Bullet")
    doc.add_heading("Education", 1)
    doc.add_paragraph("Bachelor of Commerce, University of Toronto, 2018")
    return doc


def test_a_plain_word_resume_has_nothing_to_flag(tmp_path):
    path = tmp_path / "plain.docx"
    _plain(Document()).save(str(path))
    assert _keys(path) == set()


def test_word_tables_text_boxes_columns_and_header_contact(tmp_path):
    doc = _plain(Document(), contact=False)
    table = doc.add_table(rows=2, cols=2)
    for cell in table._cells:
        cell.text = "Python SQL Excel Tableau Power BI forecasting budgeting"
    doc.sections[0].header.paragraphs[0].text = "jordan.lee@example.com | 416 555 0199"
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), "2")
    doc.sections[0]._sectPr.append(cols)
    path = tmp_path / "designed.docx"
    doc.save(str(path))
    assert {"ats_tables", "ats_columns", "ats_header"} <= _keys(path)
    assert "ats_email" not in _keys(path)  # it's there, just in the header


def test_missing_contact_and_unusual_headings(tmp_path):
    doc = Document()
    doc.add_paragraph("Jordan Lee")
    doc.add_heading("My Journey", 1)
    for _ in range(4):
        doc.add_paragraph(BULLET, style="List Bullet")
    path = tmp_path / "bare.docx"
    doc.save(str(path))
    keys = _keys(path)
    assert {"ats_email", "ats_phone", "ats_headings"} <= keys
    finding = next(f for f in ats.check(path, canadian.read(path), 1000) if f.key == "ats_headings")
    assert finding.level == canadian.SHOULD and "My Journey" in finding.examples


def test_large_files(tmp_path):
    path = tmp_path / "plain.docx"
    _plain(Document()).save(str(path))
    found = {f.key: f for f in ats.check(path, canadian.read(path), 3_200_000)}
    assert "3.2 MB" in found["ats_file_size"].title


def _pdf(path: Path, two_columns: bool) -> None:
    page = canvas.Canvas(str(path), pagesize=LETTER)
    page.setFont("Helvetica", 9)
    y = 740
    page.drawString(50, y, "Jordan Lee   jordan.lee@example.com   416 555 0199   Toronto, ON")
    for n in range(40):
        y -= 16
        if two_columns:
            page.drawString(50, y, f"Skill {n} Python SQL")
            page.drawString(260, y, "Built reporting that cut the monthly close by four days")
        else:
            page.drawString(50, y, BULLET)
            page.drawRightString(560, y, "Jan 2020")
    page.save()


def test_pdf_columns_are_told_from_a_one_column_page_with_dates_at_the_right(tmp_path):
    one, two = tmp_path / "one.pdf", tmp_path / "two.pdf"
    _pdf(one, two_columns=False)
    _pdf(two, two_columns=True)
    assert "ats_columns" not in _keys(one)
    assert "ats_columns" in _keys(two)


def test_the_sample_pdf_reads_cleanly():
    assert "ats_unreadable" not in _keys(FIXTURES / "sample_resume.pdf")


def test_every_check_has_a_basis_and_none_overclaims():
    keys = {"photo", "personal", "declaration", "references", "high_school", "dates", "spelling", "length",
            "location", "phone", "credentials", "objective", "ats_tables", "ats_columns", "ats_header",
            "ats_text_boxes", "ats_graphics", "ats_file_size", "ats_unreadable", "ats_email", "ats_phone",
            "ats_headings"}
    assert keys == set(check_sources.BASIS) == set(check_sources.RULES)
    for key, basis in check_sources.BASIS.items():
        assert basis.kind in check_sources.KIND_LABELS
        # A source-backed rule names its source; a practice says it has none.
        assert (basis.source is None) == (basis.kind == check_sources.PRACTICE) or basis.kind == check_sources.DIRECT
        assert basis.quote or basis.summary or basis.kind in (check_sources.PRACTICE, check_sources.DIRECT), key
        for claim in ("will reject", "will refuse", "always", "guarantee", "75%"):
            assert claim not in basis.impact.lower(), key
